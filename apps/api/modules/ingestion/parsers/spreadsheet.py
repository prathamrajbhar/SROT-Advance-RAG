"""Spreadsheet parsers that emit one ``sheet_row`` element per data row.

Each row's content is serialised as ``col: value; col: value; ...`` so
exact-match IDs (e.g. "INV-2024-0042") are retrievable by BM25 FTS
without being buried inside a batched blob.

Headers are included in every row's content so the chunk is
self-contained; the sheet name is carried in ``heading_path``.
"""
from __future__ import annotations

import csv
import io
import logging
from typing import List

import openpyxl

from modules.ingestion.chunker import ParsedElement

logger = logging.getLogger(__name__)

_CSV_FALLBACK_SHEET = "Sheet1"


def _row_to_content(headers: List[str], values: List[str]) -> str:
    """Serialise one row as 'col: value; col: value; ...'

    Skips columns where value is empty and returns empty string if no valid values exist.
    """
    parts: List[str] = []
    has_any_value = False
    for header, value in zip(headers, values):
        val_str = str(value).strip() if value is not None else ""
        if val_str:
            has_any_value = True
            parts.append(f"{header}: {val_str}")
    if not has_any_value:
        return ""
    # Any extra values beyond header count (rare but possible in CSV)
    for value in values[len(headers) :]:
        val_str = str(value).strip() if value is not None else ""
        if val_str:
            parts.append(val_str)
    return "; ".join(parts)


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------


def parse_csv(content_bytes: bytes) -> List[ParsedElement]:
    text = content_bytes.decode("utf-8", errors="replace")
    if not text.strip():
        return []

    delimiter = ","
    try:
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=",\t;|")
        delimiter = dialect.delimiter
    except csv.Error:
        if "\t" in text and "," not in text:
            delimiter = "\t"
        elif ";" in text and "," not in text:
            delimiter = ";"

    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = list(reader)
    if not rows:
        return []

    headers = [h.strip() for h in rows[0]]
    elements: List[ParsedElement] = []

    for row_idx, row in enumerate(rows[1:], start=2):
        values = [str(v).strip() for v in row]
        content = _row_to_content(headers, values)
        if not content.strip():
            continue
        elements.append(ParsedElement(
            kind="sheet_row",
            content=content,
            locator={
                "row_range": [row_idx, row_idx],
                "sheet": _CSV_FALLBACK_SHEET,
            },
            heading_path=[_CSV_FALLBACK_SHEET],
        ))

    logger.debug("parse_csv: %d data rows → %d sheet_row elements", len(rows) - 1, len(elements))
    return elements


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------


def parse_xlsx(content_bytes: bytes) -> List[ParsedElement]:
    wb = openpyxl.load_workbook(io.BytesIO(content_bytes), data_only=True)
    elements: List[ParsedElement] = []

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue

        headers = [
            str(h).strip() if h is not None else f"Col_{i + 1}"
            for i, h in enumerate(rows[0])
        ]

        sheet_elements: List[ParsedElement] = []
        for row_idx, row in enumerate(rows[1:], start=2):
            values = [str(v).strip() if v is not None else "" for v in row]
            content = _row_to_content(headers, values)
            if not content.strip():
                continue
            sheet_elements.append(ParsedElement(
                kind="sheet_row",
                content=content,
                locator={
                    "row_range": [row_idx, row_idx],
                    "sheet": sheet_name,
                },
                heading_path=[sheet_name],
            ))

        logger.debug(
            "parse_xlsx: sheet '%s' → %d sheet_row elements",
            sheet_name,
            len(sheet_elements),
        )
        elements.extend(sheet_elements)

    return elements
