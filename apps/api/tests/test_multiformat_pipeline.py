import io
import hashlib
import docx
import openpyxl
import pytest
from sqlalchemy import select
from core.database import async_session_factory
from core.models.factory import get_embedder
from core.qdrant import delete_project_collection, search_vectors
from core.s3 import delete_s3_prefix, upload_file_bytes
from models.auth import User
from models.chat import Conversation
from models.document import Document, IngestStatus
from models.project import Project, ProjectMember, ProjectRole
from modules.chat.stream_service import stream_chat_response
from modules.ingestion.pipeline import process_document_ingestion


@pytest.mark.asyncio
async def test_multiformat_ingestion_and_rag():
    async with async_session_factory() as db:
        user_res = await db.execute(select(User).limit(1))
        user = user_res.scalar_one_or_none()
        if not user:
            pytest.skip("No user found in database")

        project = Project(
            owner_id=user.id,
            name="Pytest Multi-Format Suite",
            description="Automated unit test for all formats",
        )
        db.add(project)
        await db.flush()

        member = ProjectMember(
            project_id=project.id, user_id=user.id, role=ProjectRole.OWNER
        )
        db.add(member)

        conv = Conversation(
            project_id=project.id, user_id=user.id, title="Test Chat"
        )
        db.add(conv)
        await db.commit()

        try:
            # 1. Prepare XLSX
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Settings"
            ws.append(["Parameter", "Value"])
            ws.append(["SessionTimeoutMinutes", "15"])
            ws.append(["MaxUploadSizeMB", "100"])
            xlsx_buf = io.BytesIO()
            wb.save(xlsx_buf)

            # 2. Prepare DOCX
            doc = docx.Document()
            doc.add_heading("Architecture Spec", level=1)
            doc.add_paragraph("SROT implements strict RBAC authorization.")
            docx_buf = io.BytesIO()
            doc.save(docx_buf)

            test_payloads = [
                (
                    "spec.md",
                    "text/markdown",
                    b"# Spec\nHybrid retrieval combines BM25 and vector search.",
                ),
                (
                    "data.csv",
                    "text/csv",
                    b"id,metric,val\n1,accuracy,0.96\n2,latency,150ms",
                ),
                (
                    "config.xlsx",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    xlsx_buf.getvalue(),
                ),
                (
                    "manual.docx",
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    docx_buf.getvalue(),
                ),
            ]

            # Ingest all formats
            for filename, mime, content in test_payloads:
                sha = hashlib.sha256(content).hexdigest()
                s3_key = f"projects/{project.id}/raw/{filename}"
                await upload_file_bytes(s3_key, content, mime)

                doc_row = Document(
                    project_id=project.id,
                    uploaded_by=user.id,
                    filename=filename,
                    mime_type=mime,
                    size_bytes=len(content),
                    sha256=sha,
                    s3_key=s3_key,
                    status=IngestStatus.QUEUED,
                )
                db.add(doc_row)
                await db.flush()

                await process_document_ingestion(db, doc_row.id, content)

                updated = (
                    await db.execute(
                        select(Document).where(Document.id == doc_row.id)
                    )
                ).scalar_one()
                assert updated.status == IngestStatus.INDEXED
                assert updated.stats is not None
                assert updated.stats.get("chunk_count", 0) > 0

            # Vector search validation
            embedder = get_embedder()
            query_vec = (await embedder.embed_texts(["hybrid retrieval"]))[0]
            pts = await search_vectors(str(project.id), query_vec, limit=5)
            assert len(pts) > 0

            # Chat streaming validation
            events = []
            async for chunk in stream_chat_response(
                conv.id,
                user.id,
                "What does hybrid retrieval combine in the spec?",
            ):
                events.append(chunk)

            raw = "".join(events)
            assert "event: token" in raw or "event: final" in raw

        finally:
            await delete_s3_prefix(f"projects/{project.id}/")
            await delete_project_collection(str(project.id))
            await db.delete(project)
            await db.commit()
